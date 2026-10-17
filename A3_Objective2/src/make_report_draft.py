"""Generate a draft 4-page A4 report (10pt, single-spaced) from outputs/. Group edits it afterwards."""
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm
from common import ROOT, OUT

doc = Document(); s = doc.sections[0]
s.page_height, s.page_width = Cm(29.7), Cm(21.0)
for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(s, m, Cm(1.8))
st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10)
st.paragraph_format.space_after = Pt(2); st.paragraph_format.line_spacing = 1.0
def H(t): p = doc.add_paragraph(); r = p.add_run(t); r.bold = True; r.font.size = Pt(11)
def P(t): doc.add_paragraph(t)
def T(csv, cols, rename=None):
    d = pd.read_csv(OUT / csv)[cols]; d = d.rename(columns=rename or {})
    t = doc.add_table(rows=1, cols=len(d.columns)); t.style = "Light Grid Accent 1"
    for i, c in enumerate(d.columns): t.rows[0].cells[i].text = str(c)
    for _, r in d.iterrows():
        cells = t.add_row().cells
        for i, v in enumerate(r): cells[i].text = f"{v:.3f}" if isinstance(v, float) else str(v)
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs: r.font.size = Pt(8)
def F(png, w=17): doc.add_picture(str(OUT / png), width=Cm(w))

p = doc.add_paragraph(); r = p.add_run("HIT140 Assessment 3 – Predicting FIFA World Cup 2026 Results with Linear Regression"); r.bold = True; r.font.size = Pt(13)
P("Darwin Group 27: Max Turner (S______), Quentin Qiu (S______), Stephen K Kenyi (S______). Code & data: <GitHub URL>")
H("1. Data and methods")
P("We compiled two datasets in Python from all 104 matches of the 2026 World Cup. Results (score after extra time, excluding shoot-outs) came from openfootball and were cross-checked against eloratings.net (all 208 team scores agree). Every explanatory variable is available before kick-off: pre-tournament FIFA ranking (Wikipedia/FIFA), Transfermarkt average player market value, average age and World Cup participations, host status, the pre-match Elo rating, and the mean goals scored/conceded in each team's previous 10 internationals. Because the form window includes pre-tournament fixtures, a team's first match uses pre-tournament statistics rather than imputed values. Skewed variables (rank, value) were log-transformed. Model 2.1 uses team1-minus-team2 differences; Model 2.2 pairs a team's own attack with its opponent's defence, so the models share no variable (limit 4). Models were fitted with statsmodels OLS and validated with repeated 5-fold cross-validation (20 repeats); competing scikit-learn algorithms were compared on identical folds.")
H("2. Linear Regression 2.1 – goal difference (n = 104)")
P("Explanatory variables: elo_diff, log_rank_diff, log_value_diff, gf10_diff, ga10_diff, age_diff, wc_exp_diff, host_diff. Goal difference averaged 0.48 (SD 2.02, range −4 to +6). In the EDA, value (r = 0.68), rank (−0.63) and Elo (0.63) differences correlated most strongly with goal difference, but also with each other (VIF up to 7).")
F("eda_2_1_heatmap.png", 9.5)
T("model_2_1_ols_coefficients.csv", ["Unnamed: 0", "coef", "std_err", "p_value", "std_coef", "vif"], {"Unnamed: 0": "variable"})
P("The model explains 53% of variance (adj. R² 0.49, F-test p < 0.001). Squad value difference dominates (b = 1.65 goals per tenfold value advantage, p < 0.001), with significant host (+0.96) and age effects; Elo and rank lose significance once value is included because they measure the same team strength (multicollinearity). Cross-validated R² is 0.38 (RMSE 1.52, MAE 1.23) versus a mean-only baseline RMSE of 2.01. Out-of-fold predictions give the correct winner in 82.5% of decisive matches; the largest errors are blow-outs (Canada 6–0 Qatar, Germany 6–0 Curaçao) and upsets (USA 3–0 Paraguay).")
F("model_2_1_pred_vs_actual.png", 15)
H("3. Linear Regression 2.2 – goals scored by a team (n = 208)")
P("Explanatory variables: elo_pre, opp_elo_pre, gf10, opp_ga10, log_avg_value, opp_log_avg_value, host, knockout. Teams scored 1.48 goals per match on average (right-skewed counts).")
T("model_2_2_ols_coefficients.csv", ["Unnamed: 0", "coef", "std_err", "p_value", "std_coef", "vif"], {"Unnamed: 0": "variable"})
P("R² = 0.31 (adj. 0.28, p < 0.001); CV R² = 0.21, RMSE 1.21, MAE 0.95 (baseline RMSE 1.39). The team's own squad value is the key driver (b = 1.18, p < 0.001); a stronger opponent (Elo) and playing at home have borderline effects (p ≈ 0.09); knockout stage has none. A single team's goal count is noisier than goal difference, which explains the lower R².")
F("model_2_2_pred_vs_actual.png", 15)
H("4. Competing models")
P("Repeated 5-fold CV RMSE (lower is better). Regularised linear models match or slightly beat OLS; tree ensembles overfit the small samples, confirming that the relationships are essentially linear.")
a = pd.read_csv(OUT / "compare_2_1.csv")[["model", "cv_r2", "cv_rmse"]].rename(columns={"cv_r2": "2.1 R²", "cv_rmse": "2.1 RMSE"})
b = pd.read_csv(OUT / "compare_2_2.csv")[["model", "cv_r2", "cv_rmse"]].rename(columns={"cv_r2": "2.2 R²", "cv_rmse": "2.2 RMSE"})
a.merge(b, on="model", how="outer").sort_values("2.2 RMSE").to_csv(OUT / "compare_both.csv", index=False)
T("compare_both.csv", ["model", "2.1 R²", "2.1 RMSE", "2.2 R²", "2.2 RMSE"])
H("5. Conclusion and limitations")
P("Pre-match information predicts goal difference moderately well and individual team goals only weakly; squad market value is the most useful single predictor. Limitations: Transfermarkt values are current rather than pre-tournament snapshots; eloratings.net is outside the listed sources; strength proxies are collinear; goal counts are non-normal (a Poisson model performed similarly). Future work could add negative-binomial or ordinal models.")
doc.save(ROOT / "report_draft.docx"); print("saved")
